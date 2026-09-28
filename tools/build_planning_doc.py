from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "交付给技术" / "创界平台生产系统建设需求与技术规划书-V1.0.md"
OUTPUT = ROOT / "交付给技术" / "创界平台生产系统建设需求与技术规划书-V1.0.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=None, bold=None, color=None, name="Arial Unicode MS"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_inline(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, color=(39, 94, 72), name="Menlo")
            run.font.size = Pt(9.5)
        else:
            run = paragraph.add_run(part)
            set_font(run)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.35)
section.right_margin = Cm(2.15)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial Unicode MS"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.35

for style_name, size, color, before, after in [
    ("Title", 28, (25, 73, 57), 0, 18),
    ("Heading 1", 18, (25, 73, 57), 18, 8),
    ("Heading 2", 14, (43, 91, 72), 14, 6),
    ("Heading 3", 11.5, (65, 92, 79), 10, 4),
]:
    style = styles[style_name]
    style.font.name = "Arial Unicode MS"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(*color)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

header = section.header.paragraphs[0]
header.text = "创界平台｜生产系统建设需求与技术规划书 V1.0"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    set_font(run, 8.5, color=(110, 125, 117))

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("内部建设资料  ·  ")
set_font(run, 8.5, color=(120, 130, 125))
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
i = 0
first_title = True
while i < len(lines):
    raw = lines[i].rstrip()
    text = raw.strip()
    if not text:
        i += 1
        continue
    if text.startswith("# "):
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, text[2:])
        first_title = False
    elif text.startswith("## "):
        doc.add_heading(text[3:], level=1)
    elif text.startswith("### "):
        doc.add_heading(text[4:], level=2)
    elif text.startswith("#### "):
        doc.add_heading(text[5:], level=3)
    elif text.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.right_indent = Cm(0.4)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        add_inline(p, text[2:])
        for run in p.runs:
            run.font.color.rgb = RGBColor(73, 92, 83)
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "EEF5F1")
        p_pr.append(shd)
    elif text.startswith("| "):
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            table_lines.append(lines[i].strip())
            i += 1
        rows = [[c.strip() for c in ln.strip("|").split("|")] for ln in table_lines]
        if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
            rows.pop(1)
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci, value in enumerate(row):
                cell = table.cell(ri, ci)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                add_inline(p, value)
                for run in p.runs:
                    set_font(run, 9, bold=(ri == 0))
                if ri == 0:
                    set_cell_shading(cell, "DCEBE3")
        set_repeat_table_header(table.rows[0])
        i -= 1
    elif re.match(r"^\d+\.\s", text):
        p = doc.add_paragraph(style="List Number")
        add_inline(p, re.sub(r"^\d+\.\s", "", text))
    elif text.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, text[2:])
    elif text == "---":
        p = doc.add_paragraph()
        p_pr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "B8CCC1")
        borders.append(bottom)
        p_pr.append(borders)
    else:
        p = doc.add_paragraph()
        if text.startswith("**版本：") or text.startswith("**日期：") or text.startswith("**用途：") or text.startswith("**依据："):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, text.replace("  ", ""))
    i += 1

doc.core_properties.title = "创界平台生产系统建设需求与技术规划书"
doc.core_properties.subject = "生产系统规划、权限、流程、数据、文件、知识库与本地大模型"
doc.core_properties.author = "创界平台"
doc.core_properties.keywords = "创界平台, 办公系统, 流程, 权限, 知识库, 本地大模型"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
